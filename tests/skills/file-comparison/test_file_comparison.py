import importlib.util
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[3]
SCRIPT_PATH = PROJECT_ROOT / "skills" / "file-comparison" / "scripts" / "file_comparison.py"


def load_module():
    spec = importlib.util.spec_from_file_location("file_comparison", SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def build_revision_docx(path: Path) -> None:
    from docx import Document

    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    document = Document()
    document.add_paragraph("第一部分  前言")
    document.add_paragraph("三、总则")
    document.add_paragraph("1、保留句。")
    document.add_paragraph("占位段落")
    document.save(path)

    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        files = {name: archive.read(name) for name in archive.namelist() if name != "word/document.xml"}

    root = ET.fromstring(document_xml)
    body = root.find(f"{{{namespace}}}body")
    assert body is not None
    paragraphs = body.findall(f"{{{namespace}}}p")
    target = paragraphs[-1]
    for child in list(target):
        target.remove(child)

    paragraph_properties = ET.SubElement(target, f"{{{namespace}}}pPr")
    ET.SubElement(paragraph_properties, f"{{{namespace}}}spacing", {f"{{{namespace}}}line": "360"})

    keep_run = ET.SubElement(target, f"{{{namespace}}}r")
    keep_text = ET.SubElement(keep_run, f"{{{namespace}}}t")
    keep_text.text = "2、原有句。"

    inserted = ET.SubElement(
        target,
        f"{{{namespace}}}ins",
        {
            f"{{{namespace}}}id": "56",
            f"{{{namespace}}}author": "tester",
            f"{{{namespace}}}date": "2020-08-18T19:44:00Z",
        },
    )
    inserted_run = ET.SubElement(inserted, f"{{{namespace}}}r")
    inserted_text = ET.SubElement(inserted_run, f"{{{namespace}}}t")
    inserted_text.text = "本基金可根据法律法规和基金合同的约定参与转融通证券出借业务。"

    deleted = ET.SubElement(
        target,
        f"{{{namespace}}}del",
        {
            f"{{{namespace}}}id": "57",
            f"{{{namespace}}}author": "tester",
            f"{{{namespace}}}date": "2020-08-18T19:45:00Z",
        },
    )
    deleted_run = ET.SubElement(deleted, f"{{{namespace}}}r")
    deleted_text = ET.SubElement(deleted_run, f"{{{namespace}}}delText")
    deleted_text.text = "这句删除内容不应回到正文。"

    updated_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", updated_xml)
        for name, content in files.items():
            archive.writestr(name, content)


def build_chinese_counting_start_docx(path: Path, *, start: int) -> None:
    from docx import Document

    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    document = Document()
    document.add_paragraph("第一部分  前言")
    document.add_paragraph("三、总则")
    document.add_paragraph("基金管理人、基金托管人在本基金合同之外披露涉及本基金的信息。")
    document.add_paragraph("本基金可根据法律法规和基金合同的约定参与转融通证券出借业务。")
    document.add_paragraph("本基金合同关于基金产品资料概要的编制、披露及更新等内容，将不晚于2020年9月1日起执行。")
    document.save(path)

    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        numbering_xml = archive.read("word/numbering.xml")
        files = {
            name: archive.read(name)
            for name in archive.namelist()
            if name not in {"word/document.xml", "word/numbering.xml"}
        }

    document_root = ET.fromstring(document_xml)
    body = document_root.find(f"{{{namespace}}}body")
    assert body is not None
    paragraphs = body.findall(f"{{{namespace}}}p")

    for paragraph in paragraphs[2:5]:
        paragraph_properties = paragraph.find(f"{{{namespace}}}pPr")
        if paragraph_properties is None:
            paragraph_properties = ET.Element(f"{{{namespace}}}pPr")
            paragraph.insert(0, paragraph_properties)
        num_pr = paragraph_properties.find(f"{{{namespace}}}numPr")
        if num_pr is None:
            num_pr = ET.Element(f"{{{namespace}}}numPr")
            paragraph_properties.insert(0, num_pr)
        ilvl = num_pr.find(f"{{{namespace}}}ilvl")
        if ilvl is None:
            ilvl = ET.SubElement(num_pr, f"{{{namespace}}}ilvl")
        ilvl.set(f"{{{namespace}}}val", "0")
        num_id = num_pr.find(f"{{{namespace}}}numId")
        if num_id is None:
            num_id = ET.SubElement(num_pr, f"{{{namespace}}}numId")
        num_id.set(f"{{{namespace}}}val", "99")

    numbering_root = ET.fromstring(numbering_xml)
    abstract = ET.SubElement(numbering_root, f"{{{namespace}}}abstractNum")
    abstract.set(f"{{{namespace}}}abstractNumId", "99")
    level = ET.SubElement(abstract, f"{{{namespace}}}lvl")
    level.set(f"{{{namespace}}}ilvl", "0")
    ET.SubElement(level, f"{{{namespace}}}start").set(f"{{{namespace}}}val", str(start))
    ET.SubElement(level, f"{{{namespace}}}numFmt").set(f"{{{namespace}}}val", "chineseCounting")
    ET.SubElement(level, f"{{{namespace}}}lvlText").set(f"{{{namespace}}}val", "%1、")

    num = ET.SubElement(numbering_root, f"{{{namespace}}}num")
    num.set(f"{{{namespace}}}numId", "99")
    ET.SubElement(num, f"{{{namespace}}}abstractNumId").set(f"{{{namespace}}}val", "99")

    updated_document_xml = ET.tostring(document_root, encoding="utf-8", xml_declaration=True)
    updated_numbering_xml = ET.tostring(numbering_root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", updated_document_xml)
        archive.writestr("word/numbering.xml", updated_numbering_xml)
        for name, content in files.items():
            archive.writestr(name, content)


def test_split_sections_handles_page_breaks_and_toc_numbers():
    module = load_module()
    text = """
封面
第一部分  前言\t1
第二部分  释义\t4

第一部分  前言
一、总则
旧内容
\f第二部分  释义
一、定义
新内容
""".strip()

    sections = module.split_sections(text)

    assert [section.number for section in sections] == ["第一部分", "第二部分"]
    assert sections[0].title == "第一部分  前言"
    assert sections[1].title == "第二部分  释义"
    assert "旧内容" in sections[0].body
    assert "新内容" in sections[1].body


def test_extract_docx_text_includes_inserted_revisions_but_ignores_deleted_revisions(tmp_path):
    load_module()
    from file_comparison.compare.extractor import extract_docx_text

    docx_path = tmp_path / "revision.docx"

    build_revision_docx(docx_path)
    text = extract_docx_text(docx_path)

    assert "本基金可根据法律法规和基金合同的约定参与转融通证券出借业务。" in text
    assert "这句删除内容不应回到正文。" not in text
    assert "2、原有句。" in text


def test_extract_docx_text_respects_word_numbering_start_value(tmp_path):
    load_module()
    from file_comparison.compare.extractor import extract_docx_text

    docx_path = tmp_path / "numbering-start.docx"

    build_chinese_counting_start_docx(docx_path, start=4)
    text = extract_docx_text(docx_path)

    assert "四、基金管理人、基金托管人在本基金合同之外披露涉及本基金的信息。" in text
    assert "五、本基金可根据法律法规和基金合同的约定参与转融通证券出借业务。" in text
    assert "六、本基金合同关于基金产品资料概要的编制、披露及更新等内容，将不晚于2020年9月1日起执行。" in text


def test_build_rows_splits_multiple_second_level_changes():
    module = load_module()
    old = module.Section(
        number="第一部分",
        title="第一部分  前言",
        body="\n".join(
            [
                "第一部分  前言",
                "三、设立方式",
                "旧段落A",
                "四、风险提示",
                "旧段落B",
            ]
        ),
    )
    new = module.Section(
        number="第一部分",
        title="第一部分  前言",
        body="\n".join(
            [
                "第一部分  前言",
                "三、设立方式",
                "新段落A",
                "四、风险提示",
                "新段落B",
            ]
        ),
    )

    rows = module.build_section_rows(old, new)

    assert len(rows) == 2
    assert [row.chapter for row in rows] == ["第一部分  前言", "第一部分  前言"]
    assert rows[0].subchapter == "三、设立方式"
    assert rows[0].old_text == "旧段落A"
    assert rows[0].new_text == "新段落A"
    assert rows[1].subchapter == "四、风险提示"
    assert rows[1].old_text == "旧段落B"
    assert rows[1].new_text == "新段落B"


def test_format_number_label_handles_word_decimal_list():
    module = load_module()

    assert module.format_number_label("decimal", "%1、", [3]) == "3、"
    assert module.format_number_label("chineseCounting", "%1、", [4]) == "四、"


def test_output_document_paths_use_two_source_file_names(tmp_path):
    module = load_module()
    pair = module.PairMatch(
        pair_id="pair-001",
        key="基金合同",
        old_path=tmp_path / "创金合信基金合同_3月.docx",
        new_path=tmp_path / "创金合信基金合同_6月.docx",
        old_label="创金合信基金合同_3月.docx",
        new_label="创金合信基金合同_6月.docx",
    )

    docx_path, doc_path = module.build_output_document_paths(pair, tmp_path, timestamp="20260501_101530")

    assert docx_path.name == "创金合信基金合同_3月 与 创金合信基金合同_6月 对照表 20260501_101530.docx"
    assert doc_path.name == "创金合信基金合同_3月 与 创金合信基金合同_6月 对照表 20260501_101530.doc"
    assert "comparison" not in docx_path.name.lower()


def test_numbered_definition_uses_short_subchapter_label():
    module = load_module()

    row = module.build_comparison_row(
        "第二部分  释义",
        "1、基金或本基金：指旧基金",
        "1、基金或本基金：指新基金",
    )

    assert row.subchapter == "1、基金或本基金："
    assert row.old_text == "1、基金或本基金：指旧基金"
    assert row.new_text == "1、基金或本基金：指新基金"


def test_build_rows_keeps_numbered_definitions_in_separate_rows():
    module = load_module()
    old = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "4、基金合同：指旧合同",
                "5、托管协议：指旧协议",
            ]
        ),
    )
    new = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "4、基金合同：指新合同",
                "5、托管协议：指新协议",
            ]
        ),
    )

    rows = module.build_section_rows(old, new)

    assert [row.subchapter for row in rows] == ["4、基金合同：", "5、托管协议："]
    assert all("……" not in row.old_text for row in rows)
    assert all("……" not in row.new_text for row in rows)


def test_build_rows_aligns_numbered_definitions_when_one_is_deleted():
    module = load_module()
    old = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "4、基金合同：指旧合同",
                "5、托管协议：指旧协议",
                "6、基金募集期：指旧募集期",
            ]
        ),
    )
    new = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "4、基金合同：指新合同",
                "5、托管协议：指新协议",
            ]
        ),
    )

    rows = module.build_section_rows(old, new)

    assert [row.subchapter for row in rows] == ["4、基金合同：", "5、托管协议：", "6、基金募集期："]
    assert rows[-1].new_text == "删除"


def test_build_rows_aligns_numbered_definitions_after_number_shift():
    module = load_module()
    old = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "7、基金份额发售公告：指旧公告",
                "8、基金产品资料概要：指旧概要",
                "9、法律法规：指旧法规",
            ]
        ),
    )
    new = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "7、基金产品资料概要：指新概要",
                "8、法律法规：指新法规",
            ]
        ),
    )

    rows = module.build_section_rows(old, new)

    assert [row.subchapter for row in rows] == [
        "7、基金份额发售公告：",
        "7、基金产品资料概要：",
        "8、法律法规：",
    ]
    assert rows[0].new_text == "删除"


def test_build_rows_skips_when_only_list_number_changes():
    module = load_module()
    old = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "8、法律法规：指中国现行有效并公布实施的法律法规",
            ]
        ),
    )
    new = module.Section(
        number="第二部分",
        title="第二部分  释义",
        body="\n".join(
            [
                "第二部分  释义",
                "7、法律法规：指中国现行有效并公布实施的法律法规",
            ]
        ),
    )

    assert module.build_section_rows(old, new) == []


def test_build_rows_omits_equal_lines_within_same_subchapter_block():
    module = load_module()
    old = module.Section(
        number="第三部分",
        title="第三部分  基金管理人",
        body="\n".join(
            [
                "第三部分  基金管理人",
                "一、基本情况",
                "设立日期：2014年7月9日",
                "批准设立机关及批准设立文号：中国证监会证监许可[2014]651号",
                "组织形式：有限责任公司",
                "注册资本：2.33亿元人民币",
                "存续期限：持续经营",
                "联系电话：0755-23838000",
            ]
        ),
    )
    new = module.Section(
        number="第三部分",
        title="第三部分  基金管理人",
        body="\n".join(
            [
                "第三部分  基金管理人",
                "一、基本情况",
                "设立日期：2014年7月9日",
                "批准设立机关及批准设立文号：中国证监会证监许可[2014]651号",
                "组织形式：股份有限公司",
                "注册资本：2.33亿元人民币",
                "存续期限：持续经营",
                "联系电话：0755-23839000",
            ]
        ),
    )

    rows = module.build_section_rows(old, new)

    assert len(rows) == 1
    assert rows[0].subchapter == "一、基本情况"
    assert rows[0].old_text == "组织形式：有限责任公司\n......\n联系电话：0755-23838000"
    assert rows[0].new_text == "组织形式：股份有限公司\n......\n联系电话：0755-23839000"


def test_remove_fully_equal_lines_marks_omitted_middle_content():
    module = load_module()
    old_text = "\n".join(["旧变化一", "中间未变化段落", "旧变化二"])
    new_text = "\n".join(["新变化一", "中间未变化段落", "新变化二"])

    old_changed, new_changed = module.remove_fully_equal_lines(old_text, new_text)

    assert old_changed == "旧变化一\n......\n旧变化二"
    assert new_changed == "新变化一\n......\n新变化二"


def test_remove_fully_equal_lines_omits_equal_numbered_clause_list():
    module = load_module()
    old_text = "\n".join(
        [
            "（七）临时报告",
            "本基金发生重大事件，有关信息披露义务人应当在2日内编制临时报告书，并登载在指定报刊和指定网站上。",
            "前款所称重大事件，是指可能对基金份额持有人权益或者基金份额的价格产生重大影响的下列事件：",
            "1、基金份额持有人大会的召开及决定的事项；",
            "2、《基金合同》终止、基金清算；",
            "3、转换基金运作方式、基金合并；",
            "21、发生涉及本基金申购、赎回事项调整或潜在影响投资者赎回等重大事项时；",
            "22、基金份额的折算；",
        ]
    )
    new_text = "\n".join(
        [
            "（七）临时报告",
            "本基金发生重大事件，有关信息披露义务人应当在2日内编制临时报告书，并登载在规定报刊和规定网站上。",
            "前款所称重大事件，是指可能对基金份额持有人权益或者基金份额的价格产生重大影响的下列事件：",
            "1、基金份额持有人大会的召开及决定的事项；",
            "2、《基金合同》终止、基金清算；",
            "3、转换基金运作方式、基金合并；",
            "21、发生涉及本基金申购、赎回事项调整或潜在影响投资者赎回等重大事项时；",
            "22、基金管理人采用摆动定价机制进行估值；",
            "23、基金份额的折算；",
        ]
    )

    old_changed, new_changed = module.remove_fully_equal_lines(old_text, new_text)

    assert old_changed == "\n".join(
        [
            "（七）临时报告",
            "本基金发生重大事件，有关信息披露义务人应当在2日内编制临时报告书，并登载在指定报刊和指定网站上。",
            "......",
        ]
    )
    assert new_changed == "\n".join(
        [
            "（七）临时报告",
            "本基金发生重大事件，有关信息披露义务人应当在2日内编制临时报告书，并登载在规定报刊和规定网站上。",
            "......",
            "22、基金管理人采用摆动定价机制进行估值；",
        ]
    )


def test_remove_fully_equal_lines_keeps_long_single_line_changes_complete():
    module = load_module()
    old_text = (
        "基金管理人可在不违反法律法规的情况下，对上述原则进行调整。"
        "基金管理人必须在新规则开始实施前依照《信息披露办法》的有关规定在指定媒介上公告。"
    )
    new_text = (
        "基金管理人可在不违反法律法规的情况下，对上述原则进行调整。"
        "基金管理人必须在新规则开始实施前依照《信息披露办法》的有关规定在规定媒介上公告。"
    )

    old_changed, new_changed = module.remove_fully_equal_lines(old_text, new_text)

    assert old_changed == old_text
    assert new_changed == new_text
    assert "......" not in old_changed
    assert "......" not in new_changed


def test_remove_fully_equal_lines_ignores_shifted_parenthesized_numbering():
    module = load_module()
    old_text = "\n".join(
        [
            "（24）基金在募集期间未能达到基金的备案条件，《基金合同》不能生效；",
            "（25）执行生效的基金份额持有人大会的决议；",
            "（26）建立并保存基金份额持有人名册；",
            "（27）法律法规及中国证监会规定的和《基金合同》约定的其他义务。",
        ]
    )
    new_text = "\n".join(
        [
            "（24）执行生效的基金份额持有人大会的决议；",
            "（25）建立并保存基金份额持有人名册；",
            "（26）法律法规及中国证监会规定的和《基金合同》约定的其他义务。",
        ]
    )

    old_changed, new_changed = module.remove_fully_equal_lines(old_text, new_text)

    assert old_changed == "（24）基金在募集期间未能达到基金的备案条件，《基金合同》不能生效；"
    assert new_changed == "删除"


def test_preprocess_sections_for_llm_skips_unchanged_sections_and_keeps_only_changed_lines():
    module = load_module()
    old_sections = [
        module.Section(
            number="第一部分",
            title="第一部分  前言",
            body="\n".join(["第一部分  前言", "一、总则", "完全一致"]),
        ),
        module.Section(
            number="第三部分",
            title="第三部分  基金管理人",
            body="\n".join(
                [
                    "第三部分  基金管理人",
                    "一、基本情况",
                    "设立日期：2014年7月9日",
                    "组织形式：有限责任公司",
                    "联系电话：0755-23838000",
                ]
            ),
        ),
    ]
    new_sections = [
        module.Section(
            number="第一部分",
            title="第一部分  前言",
            body="\n".join(["第一部分  前言", "一、总则", "完全一致"]),
        ),
        module.Section(
            number="第三部分",
            title="第三部分  基金管理人",
            body="\n".join(
                [
                    "第三部分  基金管理人",
                    "一、基本情况",
                    "设立日期：2014年7月9日",
                    "组织形式：股份有限公司",
                    "联系电话：0755-23839000",
                ]
            ),
        ),
    ]

    reduced_old, reduced_new = module.preprocess_sections_for_llm(old_sections, new_sections)

    assert [section.number for section in reduced_old] == ["第三部分"]
    assert [section.number for section in reduced_new] == ["第三部分"]
    assert "设立日期：2014年7月9日" not in reduced_old[0].body
    assert "设立日期：2014年7月9日" not in reduced_new[0].body
    assert "组织形式：有限责任公司" in reduced_old[0].body
    assert "组织形式：股份有限公司" in reduced_new[0].body


def test_preprocess_sections_for_llm_marks_added_section_for_model():
    module = load_module()
    old_sections = []
    new_sections = [
        module.Section(
            number="第四部分",
            title="第四部分  基金的历史沿革",
            body="\n".join(["第四部分  基金的历史沿革", "新增段落"]),
        )
    ]

    reduced_old, reduced_new = module.preprocess_sections_for_llm(old_sections, new_sections)

    assert reduced_old[0].body == "第四部分  基金的历史沿革\n新增"
    assert reduced_new[0].body.endswith("新增段落")


def test_build_compare_blocks_for_llm_keeps_root_siblings_in_one_block():
    module = load_module()
    old_sections = [
        module.Section(
            number="第一部分",
            title="第一部分  前言",
            body="\n".join(
                [
                    "第一部分  前言",
                    "三、创金合信中证500指数增强型发起式证券投资基金由基金管理人依照《基金法》、基金合同及其他有关规定募集。",
                    "四、基金管理人、基金托管人在本基金合同之外披露涉及本基金的信息，其内容涉及界定基金合同当事人之间权利义务关系的，如与基金合同有冲突，以基金合同为准。",
                    "五、本基金按照中国法律法规成立并运作，若基金合同的内容与届时有效的法律法规的强制性规定不一致，应当以届时有效的法律法规的规定为准。",
                    "六、本基金合同关于基金产品资料概要的编制、披露及更新等内容，将不晚于2020年9月1日起执行。",
                ]
            ),
        )
    ]
    new_sections = [
        module.Section(
            number="第一部分",
            title="第一部分  前言",
            body="\n".join(
                [
                    "第一部分  前言",
                    "三、创金合信中证500指数增强型发起式证券投资基金由基金管理人依照《基金法》、基金合同及其他有关规定募集。",
                    "四、基金管理人、基金托管人在本基金合同之外披露涉及本基金的信息，其内容涉及界定基金合同当事人之间权利义务关系的，如与基金合同有冲突，以基金合同为准。",
                    "五、本基金可根据法律法规和基金合同的约定参与转融通证券出借业务，可能存在流动性风险、市场风险和信用风险等转融通业务特有风险。",
                    "六、本基金按照中国法律法规成立并运作，若基金合同的内容与届时有效的法律法规的强制性规定不一致，应当以届时有效的法律法规的规定为准。",
                ]
            ),
        )
    ]

    blocks, summaries = module.build_compare_blocks_for_llm(old_sections, new_sections)

    assert len(blocks) == 1
    assert blocks[0].parent_path == ""
    # 去序号后「五、法律法规…」与「六、法律法规…」在邻域 5 内成对剔除，仅保留实质差异条目。
    assert [item.item_id for item in blocks[0].old_items] == ["第一部分-old-004"]
    assert [item.item_id for item in blocks[0].new_items] == ["第一部分-new-003"]
    assert [item.text for item in blocks[0].old_items] == [
        "六、本基金合同关于基金产品资料概要的编制、披露及更新等内容，将不晚于2020年9月1日起执行。",
    ]
    assert [item.text for item in blocks[0].new_items] == [
        "五、本基金可根据法律法规和基金合同的约定参与转融通证券出借业务，可能存在流动性风险、市场风险和信用风险等转融通业务特有风险。",
    ]
    assert summaries[0]["block_count"] == 1


def test_build_compare_blocks_keeps_long_numbered_root_clause_as_item_not_parent():
    module = load_module()
    old_sections = [
        module.Section(
            number="第一部分",
            title="第一部分  前言",
            body="\n".join(
                [
                    "第一部分  前言",
                    "三、创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）由基金管理人依照《基金法》、基金合同及其他有关规定募集，并经中国证券监督管理委员会注册。",
                    "中国证监会对本基金募集的注册，并不表明其对本基金的投资价值和市场前景做出实质性判断或保证，也不表明投资于本基金没有风险。",
                    "基金管理人依照恪尽职守、诚实信用、谨慎勤勉的原则管理和运用基金财产，但不保证投资于本基金一定盈利，也不保证最低收益。",
                ]
            ),
        )
    ]
    new_sections = [
        module.Section(
            number="第一部分",
            title="第一部分  前言",
            body="\n".join(
                [
                    "第一部分  前言",
                    "三、创金合信宜久来福6个月持有期混合型发起式基金中基金（FOF）由创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）转型而来。",
                    "中国证监会对本基金募集的注册，并不表明其对本基金的投资价值和市场前景做出实质性判断或保证，也不表明投资于本基金没有风险。",
                    "基金管理人依照恪尽职守、诚实信用、谨慎勤勉的原则管理和运用基金财产，但不保证投资于本基金一定盈利，也不保证最低收益。",
                ]
            ),
        )
    ]

    blocks, _summaries = module.build_compare_blocks_for_llm(old_sections, new_sections)

    assert len(blocks) == 1
    assert blocks[0].parent_path == ""
    assert blocks[0].old_items[0].text.startswith("三、创金合信宜久来福3个月持有期")
    assert blocks[0].new_items[0].text.startswith("三、创金合信宜久来福6个月持有期")
    assert "中国证监会对本基金募集的注册" in blocks[0].old_items[0].text
    assert "中国证监会对本基金募集的注册" in blocks[0].new_items[0].text


def test_build_compare_blocks_groups_parenthesized_siblings_under_parent_heading():
    module = load_module()
    old_sections = [
        module.Section(
            number="第七部分",
            title="第七部分 基金合同当事人及权利义务",
            body="\n".join(
                [
                    "第七部分 基金合同当事人及权利义务",
                    "一、基金管理人",
                    "（15）完全一致内容；",
                    "（16）办理基金认购、申购业务；",
                    "（17）继续一致内容；",
                    "（24）基金募集失败时退还认购人；",
                    "（25）执行生效的基金份额持有人大会的决议；",
                ]
            ),
        )
    ]
    new_sections = [
        module.Section(
            number="第七部分",
            title="第七部分 基金合同当事人及权利义务",
            body="\n".join(
                [
                    "第七部分 基金合同当事人及权利义务",
                    "一、基金管理人",
                    "（15）完全一致内容；",
                    "（16）办理基金申购业务；",
                    "（17）继续一致内容；",
                    "（24）执行生效的基金份额持有人大会的决议；",
                ]
            ),
        )
    ]

    blocks, summaries = module.build_compare_blocks_for_llm(old_sections, new_sections)

    assert len(blocks) == 1
    assert blocks[0].parent_path == "一、基金管理人"
    # （25）与新版（24）去序号后一致且在邻域内，成对剔除。
    assert [item.text for item in blocks[0].old_items] == [
        "（16）办理基金认购、申购业务；",
        "（24）基金募集失败时退还认购人；",
    ]
    assert [item.text for item in blocks[0].new_items] == [
        "（16）办理基金申购业务；",
    ]
    assert summaries[0]["block_count"] == 1


def test_build_compare_blocks_creates_separate_blocks_for_parenthesized_chinese_subheads():
    module = load_module()
    old_sections = [
        module.Section(
            number="第七部分",
            title="第七部分 基金合同当事人及权利义务",
            body="\n".join(
                [
                    "第七部分 基金合同当事人及权利义务",
                    "一、基金管理人",
                    "（一）基金管理人简况",
                    "名称：创金合信基金管理有限公司",
                    "法定代表人：刘学民",
                    "联系电话：0755-23838000",
                    "（二）基金管理人的权利与义务",
                    "1、根据《基金法》及其他有关规定，基金管理人的权利包括但不限于：",
                    "（17）制订和调整有关基金认购、申购、赎回等业务规则；",
                ]
            ),
        )
    ]
    new_sections = [
        module.Section(
            number="第七部分",
            title="第七部分 基金合同当事人及权利义务",
            body="\n".join(
                [
                    "第七部分 基金合同当事人及权利义务",
                    "一、基金管理人",
                    "（一）基金管理人简况",
                    "名称：创金合信基金管理有限公司",
                    "法定代表人：钱龙海",
                    "联系电话：0755-23838000",
                    "（二）基金管理人的权利与义务",
                    "1、根据《基金法》及其他有关规定，基金管理人的权利包括但不限于：",
                    "（17）制订和调整有关基金申购、赎回等业务规则；",
                ]
            ),
        )
    ]

    blocks, _summaries = module.build_compare_blocks_for_llm(old_sections, new_sections)

    assert [block.parent_path for block in blocks] == ["一、基金管理人"]
    assert [item.text for item in blocks[0].old_items] == [
        "（一）基金管理人简况\n名称：创金合信基金管理有限公司\n法定代表人：刘学民\n联系电话：0755-23838000",
        "（二）基金管理人的权利与义务\n1、根据《基金法》及其他有关规定，基金管理人的权利包括但不限于：\n（17）制订和调整有关基金认购、申购、赎回等业务规则；",
    ]
    assert [item.text for item in blocks[0].new_items] == [
        "（一）基金管理人简况\n名称：创金合信基金管理有限公司\n法定代表人：钱龙海\n联系电话：0755-23838000",
        "（二）基金管理人的权利与义务\n1、根据《基金法》及其他有关规定，基金管理人的权利包括但不限于：\n（17）制订和调整有关基金申购、赎回等业务规则；",
    ]


def test_build_compare_blocks_for_llm_skips_wholly_unchanged_chapter():
    module = load_module()
    section = module.Section(
        number="第四部分",
        title="第四部分  基金份额的发售",
        body="\n".join(
            [
                "第四部分  基金份额的发售",
                "一、基金份额的发售时间、发售方式、发售对象",
                "本基金份额发售面值为人民币1.00元。",
            ]
        ),
    )

    blocks, summaries = module.build_compare_blocks_for_llm([section], [section])

    assert blocks == []
    assert summaries[0]["block_count"] == 0
    assert summaries[0]["kept_for_llm"] is False
    assert summaries[0]["reason"] == "unchanged"


def test_build_compare_blocks_for_llm_keeps_only_changed_parent_blocks():
    module = load_module()
    old_sections = [
        module.Section(
            number="第七部分",
            title="第七部分 基金合同当事人及权利义务",
            body="\n".join(
                [
                    "第七部分 基金合同当事人及权利义务",
                    "一、基金管理人",
                    "（一）基金管理人简况",
                    "名称：创金合信基金管理有限公司",
                    "法定代表人：刘学民",
                    "（二）基金管理人的权利与义务",
                    "（17）制订和调整有关基金认购、申购、赎回等业务规则；",
                    "二、基金托管人",
                    "（一）基金托管人简况",
                    "名称：招商银行股份有限公司",
                ]
            ),
        )
    ]
    new_sections = [
        module.Section(
            number="第七部分",
            title="第七部分 基金合同当事人及权利义务",
            body="\n".join(
                [
                    "第七部分 基金合同当事人及权利义务",
                    "一、基金管理人",
                    "（一）基金管理人简况",
                    "名称：创金合信基金管理有限公司",
                    "法定代表人：钱龙海",
                    "（二）基金管理人的权利与义务",
                    "（17）制订和调整有关基金申购、赎回等业务规则；",
                    "二、基金托管人",
                    "（一）基金托管人简况",
                    "名称：招商银行股份有限公司",
                ]
            ),
        )
    ]

    blocks, summaries = module.build_compare_blocks_for_llm(old_sections, new_sections)

    assert len(blocks) == 1
    assert blocks[0].parent_path == "一、基金管理人"
    assert summaries[0]["block_count"] == 1
    assert summaries[0]["kept_for_llm"] is True


def test_build_compare_blocks_for_llm_aligns_parent_heading_after_insert():
    module = load_module()
    old_sections = [
        module.Section(
            number="第六部分",
            title="第六部分  基金份额的申购与赎回",
            body="\n".join(
                [
                    "第六部分  基金份额的申购与赎回",
                    "十六、其他",
                    "1、在对基金份额持有人利益无实质不利影响的前提下，基金管理人经与基金托管人协商一致，可对基金份额进行折算，不需召开基金份额持有人大会审议。",
                    "2、当技术条件成熟，本基金管理人在不违反法律法规且对基金份额持有人利益无实质不利影响的前提下，经与基金托管人协商一致，可根据具体情况对上述申购和赎回的安排进行补充和调整，或者开通本基金的外币申购和赎回。",
                    "3、在法律法规允许且条件具备的情况下，基金管理人可受理基金份额持有人通过中国证监会认可的交易场所或者交易方式进行份额转让的申请并由登记机构办理基金份额的过户登记。",
                ]
            ),
        )
    ]
    new_sections = [
        module.Section(
            number="第六部分",
            title="第六部分  基金份额的申购与赎回",
            body="\n".join(
                [
                    "第六部分  基金份额的申购与赎回",
                    "十六、实施侧袋机制期间本基金的申购与赎回",
                    "本基金实施侧袋机制的，本基金的申购和赎回安排详见招募说明书或相关公告。",
                    "十七、其他",
                    "1、在对基金份额持有人利益无实质不利影响的前提下，基金管理人经与基金托管人协商一致，可对基金份额进行折算，不需召开基金份额持有人大会审议。",
                    "2、当技术条件成熟，本基金管理人在不违反法律法规且对基金份额持有人利益无实质不利影响的前提下，经与基金托管人协商一致，可根据具体情况对上述申购和赎回的安排进行补充和调整，或者开通本基金的外币申购和赎回。",
                    "3、在法律法规允许且条件具备的情况下，基金管理人可受理基金份额持有人通过中国证监会认可的交易场所或者交易方式进行份额转让的申请并由登记机构办理基金份额的过户登记。",
                ]
            ),
        )
    ]

    blocks, summaries = module.build_compare_blocks_for_llm(old_sections, new_sections)

    assert [block.parent_path for block in blocks] == ["十六、实施侧袋机制期间本基金的申购与赎回"]
    assert blocks[0].old_items == ()
    assert [item.text for item in blocks[0].new_items] == [
        "本基金实施侧袋机制的，本基金的申购和赎回安排详见招募说明书或相关公告。"
    ]
    assert summaries[0]["block_count"] == 1


def test_group_compare_blocks_into_batches_uses_chapter_count_limit():
    module = load_module()
    blocks = [
        module.CompareBlock(
            block_id="第一部分-block-001",
            chapter_number="第一部分",
            chapter_title="第一部分  前言",
            parent_path="",
            old_items=(module.CompareBlockItem("old-1", "一、总则"),),
            new_items=(module.CompareBlockItem("new-1", "一、总则"),),
        ),
        module.CompareBlock(
            block_id="第二部分-block-001",
            chapter_number="第二部分",
            chapter_title="第二部分  释义",
            parent_path="",
            old_items=(module.CompareBlockItem("old-2", "1、定义A"),),
            new_items=(module.CompareBlockItem("new-2", "1、定义B"),),
        ),
        module.CompareBlock(
            block_id="第三部分-block-001",
            chapter_number="第三部分",
            chapter_title="第三部分  基金的基本情况",
            parent_path="",
            old_items=(module.CompareBlockItem("old-3", "一、基金名称\n旧名称"),),
            new_items=(module.CompareBlockItem("new-3", "一、基金名称\n新名称"),),
        ),
    ]
    batches = module.group_compare_blocks_into_batches(blocks, batch_size=2)

    assert [batch.batch_id for batch in batches] == ["batch-001", "batch-002"]
    assert batches[0].chapter_numbers == ("第一部分", "第二部分")
    assert [block.block_id for block in batches[0].compare_blocks] == [
        "第一部分-block-001",
        "第二部分-block-001",
    ]
    assert batches[1].chapter_numbers == ("第三部分",)


def test_group_compare_blocks_keeps_same_chapter_together_above_limits():
    module = load_module()
    blocks = [
        module.CompareBlock(
            block_id=f"第六部分-block-{index:03d}",
            chapter_number="第六部分",
            chapter_title="第六部分  基金份额的申购与赎回",
            parent_path=f"{index}、标题",
            old_items=(module.CompareBlockItem(f"old-{index}", f"{index}、旧条目" + "甲" * 2400),),
            new_items=(module.CompareBlockItem(f"new-{index}", f"{index}、新条目" + "乙" * 2400),),
        )
        for index in range(1, 5)
    ]

    batches = module.group_compare_blocks_into_batches(
        blocks,
        batch_size=4,
        max_batch_chars=10000,
        oversized_batch_size=2,
        max_compare_blocks_per_batch=2,
        max_compare_block_chars=10000,
    )

    assert len(batches) == 1
    assert batches[0].chapter_numbers == ("第六部分",)
    assert [block.block_id for block in batches[0].compare_blocks] == [
        "第六部分-block-001",
        "第六部分-block-002",
        "第六部分-block-003",
        "第六部分-block-004",
    ]


def test_normalize_product_name_rows_moves_name_before_preface():
    module = load_module()
    rows = [
        module.ComparisonRow(chapter="第一部分  前言", old_text="旧前言", new_text="新前言"),
        module.ComparisonRow(
            chapter="第三部分  基金的基本情况",
            subchapter="一、基金名称",
            old_text="一、基金名称\n旧产品",
            new_text="一、基金名称\n新产品",
        ),
    ]

    normalized = module.normalize_product_name_rows(rows, "旧产品", "新产品")

    assert normalized[0].chapter == "基金名称"
    assert normalized[0].old_text == "旧产品"
    assert normalized[0].new_text == "新产品"
    assert [row.chapter for row in normalized[1:]] == ["第一部分  前言"]


def test_fund_name_detection_is_optional():
    module = load_module()
    rows = [module.ComparisonRow(chapter="第一部分  前言", old_text="旧前言", new_text="新前言")]

    assert module.extract_fund_name_or_empty("第一部分  前言\n没有基金名称字段") == ""
    assert module.normalize_product_name_rows(rows, "", "") == rows


def test_apply_section_skip_rules_trims_signature_tail():
    module = load_module()
    sections = [
        module.Section(
            number="第二十四部分",
            title="第二十四部分  基金合同内容摘要",
            body="\n".join(
                [
                    "第二十四部分  基金合同内容摘要",
                    "正文条款",
                    "本页无正文，为《基金合同》的签署页。",
                    "基金管理人：某某公司（盖章）",
                ]
            ),
        )
    ]

    filtered, records = module.apply_section_skip_rules(sections, ("签署页", "签字页", "盖章页", "签章页"))

    assert filtered == [
        module.Section(
            number="第二十四部分",
            title="第二十四部分  基金合同内容摘要",
            body="第二十四部分  基金合同内容摘要\n正文条款",
        )
    ]
    assert records == [
        {
            "section_number": "第二十四部分",
            "section_title": "第二十四部分  基金合同内容摘要",
            "action": "trim_tail",
            "reason": "signature_page",
            "matched_pattern": "签署页",
            "removed_line_count": 2,
        }
    ]


def test_apply_section_skip_rules_skips_signature_only_section():
    module = load_module()
    sections = [
        module.Section(
            number="第二十五部分",
            title="第二十五部分  签署页",
            body="第二十五部分  签署页\n基金管理人：某某公司（盖章）",
        )
    ]

    filtered, records = module.apply_section_skip_rules(sections, ("签署页", "签字页", "盖章页", "签章页"))

    assert filtered == []
    assert records[0]["action"] == "skip_section"
    assert records[0]["matched_pattern"] == "签署页"


def test_old_revision_runs_strike_only_replaced_characters():
    module = load_module()

    paragraphs = module.build_old_revision_paragraphs(
        "创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）",
        "创金合信宜久来福6个月持有期混合型发起式基金中基金（FOF）",
    )

    struck = [
        text
        for paragraph in paragraphs
        for text, style in paragraph
        if style.get("strike")
    ]
    unstruck = "".join(
        text
        for paragraph in paragraphs
        for text, style in paragraph
        if not style.get("strike")
    )
    assert struck == ["3"]
    assert "创金合信宜久来福" in unstruck
    assert "个月持有期" in unstruck


def test_old_revision_runs_strike_entire_text_for_deleted_rows():
    module = load_module()

    paragraphs = module.build_old_revision_paragraphs("基金份额发售公告：旧定义", "删除")

    assert paragraphs == [[("基金份额发售公告：旧定义", {"strike": True, "color": "C00000", "size": 9.5})]]


def test_new_revision_runs_underline_only_replaced_characters():
    module = load_module()

    paragraphs = module.build_new_revision_paragraphs(
        "创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）",
        "创金合信宜久来福6个月持有期混合型发起式基金中基金（FOF）",
    )

    emphasized = [
        text
        for paragraph in paragraphs
        for text, style in paragraph
        if style.get("bold") and style.get("underline")
    ]
    normal = "".join(
        text
        for paragraph in paragraphs
        for text, style in paragraph
        if not style.get("bold") and not style.get("underline")
    )
    assert emphasized == ["6"]
    assert [
        style.get("color")
        for paragraph in paragraphs
        for text, style in paragraph
        if text == "6"
    ] == ["0070C0"]
    assert "创金合信宜久来福" in normal
    assert "个月持有期" in normal


def test_revision_runs_treat_embedded_short_equal_as_replacement():
    module = load_module()
    old_text = "依照《基金法》、基金合同及其他有关规定募集"
    new_text = "依照《基金法》、《创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）基金合同》及其他有关规定募集"

    old_paragraphs = module.build_old_revision_paragraphs(old_text, new_text)
    new_paragraphs = module.build_new_revision_paragraphs(old_text, new_text)
    struck = "".join(
        text
        for paragraph in old_paragraphs
        for text, style in paragraph
        if style.get("strike")
    )
    emphasized = "".join(
        text
        for paragraph in new_paragraphs
        for text, style in paragraph
        if style.get("bold") and style.get("color") == "0070C0"
    )

    assert "基金合同" in struck
    assert "《创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）基金合同》" in emphasized


def test_revision_runs_mark_highly_changed_short_title_as_full_replace():
    module = load_module()

    old_paragraphs = module.build_old_revision_paragraphs("第四部分  基金份额的发售", "第四部分  基金的历史沿革")
    new_paragraphs = module.build_new_revision_paragraphs("第四部分  基金份额的发售", "第四部分  基金的历史沿革")

    assert old_paragraphs == [
        [
            ("第四部分  ", {"size": 9.5}),
            ("基金份额的发售", {"strike": True, "color": "C00000", "size": 9.5}),
        ]
    ]
    assert new_paragraphs == [
        [
            ("第四部分  ", {"size": 9.5}),
            ("基金的历史沿革", {"bold": True, "underline": False, "color": "0070C0", "size": 9.5}),
        ]
    ]


def test_revision_runs_count_scattered_short_equals_as_full_replace():
    module = load_module()
    old_text = "自基金份额发售之日起最长不得超过3个月，具体发售时间见基金份额发售公告。"
    new_text = "创金合信宜久来福6个月持有期混合型发起式基金中基金（FOF）由创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）转型而来。"

    old_paragraphs = module.build_old_revision_paragraphs(old_text, new_text)
    new_paragraphs = module.build_new_revision_paragraphs(old_text, new_text)

    assert old_paragraphs == [[(old_text, {"strike": True, "color": "C00000", "size": 9.5})]]
    assert new_paragraphs == [[(new_text, {"bold": True, "underline": False, "color": "0070C0", "size": 9.5})]]


def test_revision_runs_smooth_changed_suffix_with_short_equals():
    module = load_module()
    old_text = "1、《基金合同》经基金管理人、基金托管人双方盖章以及双方法定代表人或授权代表签字或盖章并在募集结束后经基金管理人向中国证监会办理基金备案手续，并经中国证监会书面确认后生效。"
    new_text = "1、本《基金合同》由《创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）基金合同》修订而来。《基金合同》经基金管理人、基金托管人双方盖章以及双方法定代表人或授权代表签字或盖章，自202X年XX月XX日起，《创金合信宜久来福6个月持有期混合型发起式基金中基金（FOF）基金合同》正式生效，原《创金合信宜久来福3个月持有期混合型发起式基金中基金（FOF）基金合同》自同一日起失效。"

    old_paragraphs = module.build_old_revision_paragraphs(old_text, new_text)

    normal_text = "".join(
        text
        for paragraph in old_paragraphs
        for text, style in paragraph
        if not style.get("strike")
    )
    struck_text = "".join(
        text
        for paragraph in old_paragraphs
        for text, style in paragraph
        if style.get("strike")
    )

    assert "1、《基金合同》经基金管理人、基金托管人双方盖章以及双方法定代表人或授权代表签字或盖章" in normal_text
    assert struck_text == "并在募集结束后经基金管理人向中国证监会办理基金备案手续，并经中国证监会书面确认后生效。"


def test_revision_runs_mark_mostly_changed_body_as_full_replace():
    module = load_module()
    old_text = "甲" * 120 + "共同"
    new_text = "乙" * 120 + "共同"

    old_paragraphs = module.build_old_revision_paragraphs(old_text, new_text)
    new_paragraphs = module.build_new_revision_paragraphs(old_text, new_text)

    assert old_paragraphs == [[(old_text, {"strike": True, "color": "C00000", "size": 9.5})]]
    assert new_paragraphs == [[(new_text, {"bold": True, "underline": False, "color": "0070C0", "size": 9.5})]]


def test_new_revision_runs_underline_entire_text_for_added_rows():
    module = load_module()

    paragraphs = module.build_new_revision_paragraphs("新增", "新增章节内容\n第二行")

    assert paragraphs == [
        [("新增章节内容", {"bold": True, "underline": True, "color": "0070C0", "size": 9.5})],
        [("第二行", {"bold": True, "underline": True, "color": "0070C0", "size": 9.5})],
    ]


def test_new_revision_runs_do_not_underline_long_added_paragraphs():
    module = load_module()
    long_text = "这是一段较长的新增说明文字，用于避免整段文字都出现下划线导致阅读困难。"

    paragraphs = module.build_new_revision_paragraphs("新增", long_text)

    assert paragraphs == [[(long_text, {"bold": True, "underline": False, "color": "0070C0", "size": 9.5})]]


def test_new_revision_runs_do_not_underline_when_most_of_line_changed():
    module = load_module()

    paragraphs = module.build_new_revision_paragraphs(
        "（26）基金合同约定的其他义务。",
        "（26）法律法规及中国证监会规定的和《基金合同》约定的其他义务。",
    )

    changed_styles = [
        style
        for paragraph in paragraphs
        for text, style in paragraph
        if style.get("color") == "0070C0"
    ]
    assert changed_styles
    assert all(not style.get("underline") for style in changed_styles)


def test_display_text_includes_subchapter_when_body_does_not():
    module = load_module()
    row = module.ComparisonRow("第一部分  前言", "旧段落A", "新段落A", "三、设立方式")

    assert module.display_text_with_subchapter(row.old_text, row.subchapter) == "三、设立方式\n旧段落A"
    assert module.display_text_with_subchapter("三、设立方式\n旧段落A", row.subchapter) == "三、设立方式\n旧段落A"
    assert module.display_text_with_subchapter("（16）办理基金申购业务；", "一、基金管理人") == "一、基金管理人\n（16）办理基金申购业务；"
    assert module.display_text_with_subchapter(
        "8、基金产品资料概要：指旧概要",
        "7、基金产品资料概要：",
    ) == "8、基金产品资料概要：指旧概要"
    assert module.display_text_with_subchapter(
        "一、基金管理人\n（一）基金管理人简况\n法定代表人：刘学民",
        "一、基金管理人\n（一）基金管理人简况",
    ) == "一、基金管理人\n（一）基金管理人简况\n法定代表人：刘学民"
    assert module.display_text_with_subchapter(
        "五、公开披露的基金信息\n（四）基金净值信息\n旧正文",
        "五、公开披露的基金信息\n（二）基金净值信息",
    ) == "五、公开披露的基金信息\n（四）基金净值信息\n旧正文"


def test_write_docx_hides_subchapter_column_and_merges_same_chapter(tmp_path):
    pytest.importorskip("docx")
    module = load_module()
    output = tmp_path / "comparison.docx"
    rows = [
        module.ComparisonRow("第一部分  前言", "旧A", "新A", "三、设立方式"),
        module.ComparisonRow("第一部分  前言", "旧B", "新B", "四、风险提示"),
        module.ComparisonRow("第二部分  释义", "1、基金或本基金：指旧", "1、基金或本基金：指新", "1、基金或本基金："),
    ]

    module.write_docx(rows, output, "旧基金", "新基金")

    from docx import Document
    from docx.oxml.ns import qn

    document = Document(output)
    table = document.tables[0]
    assert len(table.columns) == 3
    assert table.rows[0].cells[1].text == "原《旧基金》版本"
    assert table.rows[0].cells[2].text == "修订后《新基金》版本"
    assert "二级内容" not in "\n".join(cell.text for row in table.rows for cell in row.cells)
    assert table.rows[2].cells[1].text.startswith("三、设立方式")

    first_data_cell_pr = table._tbl.tr_lst[2].tc_lst[0].tcPr
    second_data_cell_pr = table._tbl.tr_lst[3].tc_lst[0].tcPr
    assert first_data_cell_pr.find(qn("w:vMerge")).get(qn("w:val")) == "restart"
    assert second_data_cell_pr.find(qn("w:vMerge")).get(qn("w:val")) is None


def test_write_docx_uses_compact_margins_and_chapter_column(tmp_path):
    pytest.importorskip("docx")
    module = load_module()
    output = tmp_path / "comparison.docx"

    module.write_docx(
        [module.ComparisonRow("第一部分  前言", "旧内容", "新内容")],
        output,
        "旧基金",
        "新基金",
    )

    from docx import Document

    document = Document(output)
    table = document.tables[0]
    assert round(document.sections[0].left_margin.cm, 1) == 1.0
    assert round(document.sections[0].right_margin.cm, 1) == 1.0
    assert round(table.columns[0].width.cm, 1) == 1.4


def _chunking_module():
    skill_src = PROJECT_ROOT / "skills" / "file-comparison" / "src"
    path = str(skill_src)
    if path not in sys.path:
        sys.path.insert(0, path)
    from file_comparison.compare import chunking

    return chunking


def test_filter_compare_block_items_by_stripped_numbering_identity_pairs_within_neighbor_window():
    chunking = _chunking_module()
    from file_comparison.compare.models import CompareBlockItem

    old_items = (
        CompareBlockItem("old-a", "4、基金份额：指同一正文。"),
        CompareBlockItem("old-b", "5、另一条：旧。"),
    )
    new_items = (
        CompareBlockItem("new-a", "6、基金份额：指同一正文。"),
        CompareBlockItem("new-b", "7、另一条：新。"),
    )
    kept_old, kept_new = chunking.filter_compare_block_items_by_stripped_numbering_identity(old_items, new_items)
    assert [item.text for item in kept_old] == ["5、另一条：旧。"]
    assert [item.text for item in kept_new] == ["7、另一条：新。"]


def test_filter_compare_block_items_by_stripped_numbering_identity_pairs_misaligned_within_window():
    """邻域内下标不要求一致：去序号后相同且 |i-j|<=5 即可成对剔除。"""
    chunking = _chunking_module()
    from file_comparison.compare.models import CompareBlockItem

    old_items = (
        CompareBlockItem("old-1", "五、本基金按照法律法规成立并运作。"),
        CompareBlockItem("old-2", "六、本基金合同关于资料概要的编制。"),
    )
    new_items = (
        CompareBlockItem("new-1", "五、本基金可参与转融通业务。"),
        CompareBlockItem("new-2", "六、本基金按照法律法规成立并运作。"),
    )
    kept_old, kept_new = chunking.filter_compare_block_items_by_stripped_numbering_identity(old_items, new_items)
    assert [item.text for item in kept_old] == ["六、本基金合同关于资料概要的编制。"]
    assert [item.text for item in kept_new] == ["五、本基金可参与转融通业务。"]


def test_filter_compare_block_items_by_stripped_numbering_identity_skips_pair_beyond_neighbor_window():
    chunking = _chunking_module()
    from file_comparison.compare.models import CompareBlockItem

    old_items = (
        CompareBlockItem("o1", "1、同体一句。"),
        CompareBlockItem("o2", "2、填老B。"),
        CompareBlockItem("o3", "3、填老C。"),
        CompareBlockItem("o4", "4、填老D。"),
        CompareBlockItem("o5", "5、填老E。"),
        CompareBlockItem("o6", "6、填老F。"),
        CompareBlockItem("o7", "7、填老G。"),
    )
    new_items = (
        CompareBlockItem("n1", "1、填新1。"),
        CompareBlockItem("n2", "2、填新2。"),
        CompareBlockItem("n3", "3、填新3。"),
        CompareBlockItem("n4", "4、填新4。"),
        CompareBlockItem("n5", "5、填新5。"),
        CompareBlockItem("n6", "6、填新6。"),
        CompareBlockItem("n7", "9、同体一句。"),
    )
    kept_old, kept_new = chunking.filter_compare_block_items_by_stripped_numbering_identity(old_items, new_items)
    assert kept_old == old_items and kept_new == new_items


def test_filter_compare_block_items_by_stripped_numbering_identity_keeps_different_body():
    chunking = _chunking_module()
    from file_comparison.compare.models import CompareBlockItem

    old_items = (CompareBlockItem("old-1", "1、甲条款"),)
    new_items = (CompareBlockItem("new-1", "1、乙条款"),)
    kept_old, kept_new = chunking.filter_compare_block_items_by_stripped_numbering_identity(old_items, new_items)
    assert kept_old == old_items and kept_new == new_items
